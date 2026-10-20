"""Escritura en tiempo real del Excel de salida indicado en `plantilla.csv`.

El contador escribe `#excel=<ruta absoluta>` y la app agrega ahí las facturas
nuevas después de cada escaneo, en una hoja llamada `Facturas`. El CSV descargable
no cambia: este es un canal aparte.

El avance se guarda en `.excel_estado.json`, en la raíz del proyecto, con un cursor
por canal (`ultimo_id_excel` y `ultimo_id_avisos`) para que un fallo de escritura
no duplique filas al reintentar en el siguiente escaneo.
"""

from __future__ import annotations

import hashlib
import json
import logging
import os
import shutil
import tempfile
import threading
from collections.abc import Sequence
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.config import RAIZ
from app.db import queries
from app.export.avisos import agregar_avisos, escribir_avisos
from app.export.csv import valores_por_frase
from app.export.plantilla import Plantilla
from app.models.invoice import Columna, Formato
from app.ocr.extract import normalizar_frase

HOJA = "Facturas"
AVISOS = "avisos.csv"
ESTADO = RAIZ / ".excel_estado.json"

_registro = logging.getLogger(__name__)
_candado = threading.Lock()


class _HojaAusente(Exception):
    """El libro existe pero perdió la hoja `Facturas`: hay que rehacerlo."""


@dataclass
class _Estado:
    """Hasta qué factura llegó cada canal y con qué encabezados se escribió."""

    ruta: str
    hash_encabezados: str
    ultimo_id_excel: int = 0
    ultimo_id_avisos: int = 0


def _leer_estado(ruta: Path) -> _Estado | None:
    """Devuelve el estado guardado, o `None` si falta o está corrupto."""
    try:
        datos = json.loads(ruta.read_text(encoding="utf-8"))
        return _Estado(
            ruta=str(datos["ruta"]),
            hash_encabezados=str(datos["hash_encabezados"]),
            ultimo_id_excel=int(datos["ultimo_id_excel"]),
            ultimo_id_avisos=int(datos["ultimo_id_avisos"]),
        )
    except (OSError, ValueError, KeyError, TypeError):
        return None


def _escribir_estado(ruta: Path, estado: _Estado) -> None:
    """Guarda el estado de forma atómica."""
    carga = {
        "ruta": estado.ruta,
        "hash_encabezados": estado.hash_encabezados,
        "ultimo_id_excel": estado.ultimo_id_excel,
        "ultimo_id_avisos": estado.ultimo_id_avisos,
    }
    descriptor, temporal = tempfile.mkstemp(
        prefix=".estado-", suffix=".json", dir=ruta.parent
    )
    os.close(descriptor)
    try:
        Path(temporal).write_text(
            json.dumps(carga, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        os.replace(temporal, ruta)
    finally:
        Path(temporal).unlink(missing_ok=True)


def _hash_encabezados(columnas: Sequence[Columna]) -> str:
    """Identifica la forma de la tabla: si cambia, hay que reescribirla."""
    texto = "\x1f".join(columna.encabezado for columna in columnas)
    return hashlib.sha256(texto.encode("utf-8")).hexdigest()


def _claves(columnas: Sequence[Columna], acentos: bool) -> list[str]:
    return [normalizar_frase(columna.frase, acentos) for columna in columnas]


def _celda(valor: Any) -> Any:
    """Los números van como números para que Excel pueda sumarlos."""
    if valor is None:
        return None
    if isinstance(valor, bool):
        return valor
    if isinstance(valor, int | float):
        return round(float(valor), 2)
    return str(valor)


def _fila(
    factura: dict[str, Any], claves: Sequence[str], acentos: bool
) -> list[Any]:
    guardados = valores_por_frase(factura.get("values") or [], acentos)
    return [_celda(guardados.get(clave)) for clave in claves]


def _rellenar(
    hoja: Worksheet,
    columnas: Sequence[Columna],
    formato: Formato,
    facturas: Sequence[dict[str, Any]],
) -> None:
    hoja.append([columna.encabezado for columna in columnas])
    claves = _claves(columnas, formato.acentos)
    for factura in facturas:
        hoja.append(_fila(factura, claves, formato.acentos))


def _respaldo(ruta: Path) -> None:
    """Copia el archivo a `*.bak` antes de reescribirlo."""
    if ruta.exists():
        shutil.copy2(ruta, ruta.with_name(ruta.name + ".bak"))


def _guardar(libro, destino: Path) -> None:
    """Guarda primero en un temporal y luego reemplaza, para no dejarlo a medias."""
    descriptor, temporal = tempfile.mkstemp(
        prefix=".facturas-", suffix=".xlsx", dir=destino.parent
    )
    os.close(descriptor)
    try:
        libro.save(temporal)
        os.replace(temporal, destino)
    finally:
        Path(temporal).unlink(missing_ok=True)


def _escribir_completo(
    destino: Path,
    plantilla: Plantilla,
    facturas: Sequence[dict[str, Any]],
) -> None:
    """Crea el libro o reemplaza solo la hoja `Facturas`, conservando las demás."""
    if destino.exists():
        libro = load_workbook(destino)
        if HOJA in libro.sheetnames:
            del libro[HOJA]
        hoja = libro.create_sheet(HOJA)
    else:
        libro = Workbook()
        hoja = libro.active
        hoja.title = HOJA

    try:
        _rellenar(hoja, plantilla.columnas, plantilla.formato, facturas)
        _guardar(libro, destino)
    finally:
        libro.close()


def _agregar(
    destino: Path,
    plantilla: Plantilla,
    facturas: Sequence[dict[str, Any]],
) -> None:
    """Agrega filas al final de la hoja existente."""
    libro = load_workbook(destino)
    try:
        if HOJA not in libro.sheetnames:
            raise _HojaAusente
        hoja = libro[HOJA]
        claves = _claves(plantilla.columnas, plantilla.formato.acentos)
        for factura in facturas:
            hoja.append(_fila(factura, claves, plantilla.formato.acentos))
        _guardar(libro, destino)
    finally:
        libro.close()


def _rehacer_todo(
    destino: Path,
    plantilla: Plantilla,
    facturas: Sequence[dict[str, Any]],
    ruta_estado: Path,
    hash_actual: str,
) -> None:
    """Reescribe libro y avisos desde cero, con respaldo de ambos."""
    _respaldo(destino)
    _escribir_completo(destino, plantilla, facturas)

    ruta_avisos = destino.parent / AVISOS
    _respaldo(ruta_avisos)
    escribir_avisos(ruta_avisos, facturas)

    ultimo = max((int(factura.get("id") or 0) for factura in facturas), default=0)
    _escribir_estado(
        ruta_estado, _Estado(str(destino), hash_actual, ultimo, ultimo)
    )


def _sincronizar(
    plantilla: Plantilla,
    ruta_estado: Path,
    facturas: Sequence[dict[str, Any]] | None,
) -> None:
    destino = plantilla.salida
    if destino is None:
        return

    if not destino.is_absolute() or destino.suffix.lower() != ".xlsx":
        _registro.warning("Ruta de Excel inválida en plantilla.csv: %s", destino)
        return

    carpeta = destino.parent
    if not carpeta.is_dir():
        _registro.warning("No existe la carpeta del Excel de salida: %s", carpeta)
        return

    datos = list(facturas) if facturas is not None else queries.listar_para_exportar()
    datos.sort(key=lambda factura: int(factura.get("id") or 0))

    estado = _leer_estado(ruta_estado)
    hash_actual = _hash_encabezados(plantilla.columnas)

    rehacer = (
        estado is None
        or not destino.exists()
        or estado.ruta != str(destino)
        or estado.hash_encabezados != hash_actual
    )
    if rehacer:
        _rehacer_todo(destino, plantilla, datos, ruta_estado, hash_actual)
        return

    nuevos = [
        factura
        for factura in datos
        if int(factura.get("id") or 0) > estado.ultimo_id_excel
    ]
    nuevos_avisos = [
        factura
        for factura in datos
        if int(factura.get("id") or 0) > estado.ultimo_id_avisos
    ]
    if not nuevos and not nuevos_avisos:
        return

    if nuevos:
        try:
            _agregar(destino, plantilla, nuevos)
        except _HojaAusente:
            _rehacer_todo(destino, plantilla, datos, ruta_estado, hash_actual)
            return
        estado.ultimo_id_excel = max(int(factura["id"]) for factura in nuevos)
        _escribir_estado(ruta_estado, estado)

    if nuevos_avisos:
        agregar_avisos(carpeta / AVISOS, nuevos_avisos)
        estado.ultimo_id_avisos = max(
            int(factura["id"]) for factura in nuevos_avisos
        )
        _escribir_estado(ruta_estado, estado)


def sincronizar(
    plantilla: Plantilla,
    *,
    ruta_estado: Path | None = None,
    facturas: Sequence[dict[str, Any]] | None = None,
) -> None:
    """Agrega al Excel de salida las facturas nuevas, sin romper el escaneo.

    Si `plantilla.salida` es `None` no hace nada. Un fallo de escritura (por
    ejemplo, el libro abierto en Excel) se registra y no avanza el cursor, para
    reintentarlo en el siguiente escaneo.
    """
    if plantilla.salida is None:
        return

    with _candado:
        try:
            _sincronizar(
                plantilla,
                ESTADO if ruta_estado is None else ruta_estado,
                facturas,
            )
        except Exception:
            _registro.exception(
                "No se pudo actualizar el Excel de salida (%s)", plantilla.salida
            )


__all__ = ["ESTADO", "HOJA", "sincronizar"]
