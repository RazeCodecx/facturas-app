"""Tests de la escritura en tiempo real del Excel de salida."""

from __future__ import annotations

import csv
import io
import json
from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook, load_workbook

from app.export import excel as modulo_excel
from app.export.excel import HOJA, sincronizar
from app.export.plantilla import Plantilla
from app.models.invoice import Columna, Formato


def _plantilla(destino: Path | None, encabezado: str = "Base imponible") -> Plantilla:
    return Plantilla([Columna("Total", encabezado, 1)], Formato(), destino)


def _factura(
    invoice_id: int,
    valor: Any = 116.0,
    motivo: str | None = None,
) -> dict[str, Any]:
    return {
        "id": invoice_id,
        "review_status": "needs_review" if motivo else "ok",
        "review_reason": motivo,
        "values": [
            {
                "frase": "Total",
                "encabezado": "Base imponible",
                "value": valor,
                "confidence": 90.0,
            }
        ],
    }


def _sincronizar(
    tmp_path: Path,
    destino: Path,
    facturas: list[dict[str, Any]],
    plantilla: Plantilla | None = None,
) -> Path:
    estado = tmp_path / ".estado.json"
    sincronizar(
        plantilla or _plantilla(destino),
        ruta_estado=estado,
        facturas=facturas,
    )
    return estado


def _avisos(tmp_path: Path) -> list[list[str]]:
    texto = (tmp_path / "avisos.csv").read_text(encoding="utf-8-sig")
    return list(csv.reader(io.StringIO(texto)))


def test_crea_libro_con_cabecera_y_numeros(tmp_path: Path) -> None:
    destino = tmp_path / "facturas.xlsx"

    estado = _sincronizar(tmp_path, destino, [_factura(1)])

    libro = load_workbook(destino)
    hoja = libro[HOJA]
    assert hoja.max_row == 2
    assert hoja["A1"].value == "Base imponible"
    assert hoja["A2"].value == 116.0
    assert isinstance(hoja["A2"].value, int | float)
    libro.close()

    guardado = json.loads(estado.read_text(encoding="utf-8"))
    assert guardado["ultimo_id_excel"] == 1
    assert guardado["ultimo_id_avisos"] == 1
    assert _avisos(tmp_path) == [["invoice_id", "frase", "columna", "motivo"]]


def test_agrega_solo_las_nuevas(tmp_path: Path) -> None:
    destino = tmp_path / "facturas.xlsx"

    _sincronizar(tmp_path, destino, [_factura(1, 100.0)])
    _sincronizar(
        tmp_path,
        destino,
        [_factura(1, 100.0), _factura(2, 200.0)],
    )

    libro = load_workbook(destino)
    hoja = libro[HOJA]
    assert hoja.max_row == 3
    assert [hoja["A2"].value, hoja["A3"].value] == [100.0, 200.0]
    libro.close()


def test_no_duplica_al_repetir_el_mismo_lote(tmp_path: Path) -> None:
    destino = tmp_path / "facturas.xlsx"
    facturas = [_factura(1)]

    _sincronizar(tmp_path, destino, facturas)
    _sincronizar(tmp_path, destino, facturas)

    libro = load_workbook(destino)
    assert libro[HOJA].max_row == 2
    libro.close()


def test_cambio_de_encabezados_reescribe_y_respalda(tmp_path: Path) -> None:
    destino = tmp_path / "facturas.xlsx"

    _sincronizar(tmp_path, destino, [_factura(1)], _plantilla(destino, "Base"))
    _sincronizar(
        tmp_path,
        destino,
        [_factura(1), _factura(2)],
        _plantilla(destino, "Total"),
    )

    assert (tmp_path / "facturas.xlsx.bak").exists()
    libro = load_workbook(destino)
    hoja = libro[HOJA]
    assert hoja["A1"].value == "Total"
    assert hoja.max_row == 3
    libro.close()


def test_conserva_las_otras_hojas(tmp_path: Path) -> None:
    destino = tmp_path / "facturas.xlsx"
    libro = Workbook()
    libro.active.title = "Datos"
    libro["Datos"]["A1"] = "no tocar"
    libro.save(destino)
    libro.close()

    _sincronizar(tmp_path, destino, [_factura(1)])

    libro = load_workbook(destino)
    assert libro["Datos"]["A1"].value == "no tocar"
    assert HOJA in libro.sheetnames
    libro.close()


def test_estado_ausente_con_libro_existente_hace_respaldo(tmp_path: Path) -> None:
    destino = tmp_path / "facturas.xlsx"
    libro = Workbook()
    hoja = libro.active
    hoja.title = HOJA
    hoja.append(["viejo"])
    libro.save(destino)
    libro.close()

    _sincronizar(tmp_path, destino, [_factura(1)])

    assert (tmp_path / "facturas.xlsx.bak").exists()
    libro = load_workbook(destino)
    hoja = libro[HOJA]
    assert hoja["A1"].value == "Base imponible"
    assert hoja.max_row == 2
    libro.close()


def test_archivo_bloqueado_no_avanza_el_cursor(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    destino = tmp_path / "facturas.xlsx"
    estado = tmp_path / ".estado.json"

    def _bloqueado(libro, ruta) -> None:
        raise PermissionError("abierto en Excel")

    monkeypatch.setattr(modulo_excel, "_guardar", _bloqueado)
    _sincronizar(tmp_path, destino, [_factura(1)])

    assert not destino.exists()
    assert not estado.exists()

    monkeypatch.undo()
    _sincronizar(tmp_path, destino, [_factura(1)])

    assert destino.exists()
    assert json.loads(estado.read_text(encoding="utf-8"))["ultimo_id_excel"] == 1


def test_sin_ruta_no_escribe_nada(tmp_path: Path) -> None:
    sincronizar(
        _plantilla(None),
        ruta_estado=tmp_path / ".estado.json",
        facturas=[_factura(1)],
    )

    assert list(tmp_path.iterdir()) == []


def test_no_escribe_datos_internos(tmp_path: Path) -> None:
    destino = tmp_path / "facturas.xlsx"
    motivo = 'no_match:Total (columna "Base imponible")'

    _sincronizar(tmp_path, destino, [_factura(1, motivo=motivo)])

    libro = load_workbook(destino)
    valores = [
        celda.value
        for fila in libro[HOJA].iter_rows()
        for celda in fila
        if celda.value is not None
    ]
    libro.close()

    for prohibido in ("Total", "needs_review", "review_reason", motivo):
        assert prohibido not in valores


def test_avisos_se_agregan_una_sola_vez(tmp_path: Path) -> None:
    destino = tmp_path / "facturas.xlsx"
    motivo = 'no_match:Total (columna "Base imponible")'

    _sincronizar(tmp_path, destino, [_factura(1, motivo=motivo)])
    _sincronizar(
        tmp_path,
        destino,
        [_factura(1, motivo=motivo), _factura(2)],
    )
    _sincronizar(
        tmp_path,
        destino,
        [_factura(1, motivo=motivo), _factura(2)],
    )

    filas = _avisos(tmp_path)
    assert filas[0] == ["invoice_id", "frase", "columna", "motivo"]
    assert filas[1] == ["1", "Total", "Base imponible", "no_match"]
    assert len(filas) == 2


def test_hoja_borrada_se_rehace_completa(tmp_path: Path) -> None:
    destino = tmp_path / "facturas.xlsx"

    _sincronizar(tmp_path, destino, [_factura(1)])
    libro = load_workbook(destino)
    libro.create_sheet("Datos")
    del libro[HOJA]
    libro.save(destino)
    libro.close()

    _sincronizar(tmp_path, destino, [_factura(1), _factura(2)])

    libro = load_workbook(destino)
    assert libro[HOJA].max_row == 3
    libro.close()
